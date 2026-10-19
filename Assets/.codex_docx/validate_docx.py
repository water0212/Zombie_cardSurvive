import sys
import zipfile

from docx import Document


path = sys.argv[1]
document = Document(path)
with zipfile.ZipFile(path) as archive:
    print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} sections={len(document.sections)}")
    print(f"zip_test={archive.testzip()}")

print("headings:")
for paragraph in document.paragraphs:
    if paragraph.style.name.startswith("Heading"):
        print(f"- {paragraph.style.name}: {paragraph.text}")
