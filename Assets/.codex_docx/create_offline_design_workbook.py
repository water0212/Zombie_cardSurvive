from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"C:\Users\water\Desktop\GameDev\Zombie_cardSurvive\Assets\Docs\GameDesign\殭屍生存卡牌_離線設計決策工作簿.docx")
FONT = "Microsoft JhengHei"
INK = "252525"
GRAY = "F1F1EF"
PALE = "F7F7F5"
HEADER = "414141"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, size=None, bold=None, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    set_paragraph_keep(p, keep_next=True)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        style_run(p.add_run(bold_lead), bold=True)
        style_run(p.add_run(text[len(bold_lead):]))
    else:
        style_run(p.add_run(text))
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        style_run(p.add_run(item))


def add_checklist(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        style_run(p.add_run("☐  " + item))


def add_prompt(doc, label, lines=2):
    p = doc.add_paragraph()
    style_run(p.add_run(label), bold=True)
    p.paragraph_format.space_after = Pt(3)
    for _ in range(lines):
        q = doc.add_paragraph("________________________________________________________________________________")
        q.paragraph_format.space_after = Pt(1)
        for run in q.runs:
            style_run(run, size=9, color="8A8A8A")


def add_two_col_table(doc, rows, widths=(5.0, 11.5), header=None):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if header:
        cells = table.add_row().cells
        for idx, value in enumerate(header):
            cells[idx].width = Cm(widths[idx])
            set_cell_shading(cells[idx], HEADER)
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_run(p.add_run(value), bold=True, color="FFFFFF")
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_repeat_table_header(table.rows[0])
    for i, (left, right) in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate((left, right)):
            cells[idx].width = Cm(widths[idx])
            set_cell_shading(cells[idx], PALE if i % 2 else "FFFFFF")
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx])
            p = cells[idx].paragraphs[0]
            style_run(p.add_run(value))
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title_style = styles["Title"]
    title_style.font.name = FONT
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title_style.font.size = Pt(25)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        s = styles[name]
        s.font.name = FONT
        s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        s.paragraph_format.space_after = Pt(5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("殭屍生存卡牌離線設計決策工作簿")
    title.paragraph_format.space_before = Pt(40)
    title.paragraph_format.space_after = Pt(14)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(subtitle.add_run("沒有電腦時也能完成的規則 內容與參數決策"), size=13, color="555555")
    subtitle.paragraph_format.space_after = Pt(28)

    add_body(doc, "這份工作簿用來記錄不依賴 Unity 畫面或程式測試的設計決定。你可以直接在手機或紙上回答，再交給開發流程轉成 GDD、Backlog 與資料 Asset。")
    add_body(doc, "優先處理規則、內容、相對強弱與玩家體驗。精確座標、動畫手感和最終平衡應等到能操作遊戲時再調整。")

    add_heading(doc, "快速使用方式", 1)
    add_checklist(doc, [
        "每次只挑一個主題，不必一次填完整份文件。",
        "先寫目的與相對關係，再填暫定數字。",
        "不確定時可以留下兩個候選方案，並寫下各自的優缺點。",
        "標記哪些內容已確認、哪些只是待測假設。",
    ])

    add_two_col_table(doc, [
        ("適合離線決定", "玩法規則、建築定位、卡牌效果、事件選項、解鎖順序、體驗目標、數值範圍"),
        ("需要電腦驗證", "UI 座標、字級、動畫時間、輸入手感、效能、最終平衡與提示可讀性"),
    ], header=("決策類型", "內容"))

    add_heading(doc, "目前最值得先決定的問題", 1)
    add_body(doc, "先完成每棟建築的定位、等級解鎖與可製作內容。這會同時影響卡牌來源、資源用途、基地成長與情報中心功能。")

    add_page_break(doc)
    add_heading(doc, "建築規則", 1)
    add_body(doc, "建築規則要先回答玩家何時能使用功能，以及升級是否會中斷既有操作。這些規則會直接影響建築面板和 Runtime 的判斷。")
    add_checklist(doc, [
        "建築升級期間，原本的製作列是否繼續生產。",
        "建築升級期間，食物計畫是否仍可調整。",
        "建築升級期間，卡牌強化是否仍可使用。",
        "情報中心未建造前，食物計畫是完全鎖定或保留基礎版本。",
        "建築是否會受事件影響、暫停、損壞或降低效率。",
        "公開但未解鎖的功能是否允許玩家預覽。",
    ])
    add_prompt(doc, "已確認的建築共通規則", 4)
    add_prompt(doc, "仍需測試的建築規則", 3)

    add_heading(doc, "單棟建築資料卡", 2)
    add_two_col_table(doc, [
        ("建築名稱", ""),
        ("設計定位", ""),
        ("最高等級", ""),
        ("玩家為什麼要建造", ""),
        ("主要消耗資源", ""),
        ("建築受到的限制", ""),
    ], widths=(5.2, 11.3))

    add_heading(doc, "建築等級規劃", 2)
    add_two_col_table(doc, [
        ("等級一", "建造成本、施工需求、解鎖功能、被動效果"),
        ("等級二", "升級成本、新增功能、數值提升"),
        ("等級三", "升級成本、新增功能、玩法變化"),
        ("最高等級", "最終收益，以及為何值得完成"),
    ], widths=(4.2, 12.3))

    add_page_break(doc)
    add_heading(doc, "卡牌與食譜", 1)
    add_body(doc, "卡牌設計先確定用途、代價與使用情境。數字可以先暫定，但每張卡都應有清楚的選擇理由。")
    add_two_col_table(doc, [
        ("卡牌名稱", ""),
        ("卡牌類型", "食物、資源、戰鬥、探索或特殊"),
        ("主要功能", ""),
        ("能量費用", ""),
        ("使用次數", ""),
        ("製作建築", ""),
        ("製作成本", ""),
        ("製造力需求", ""),
        ("升級版本", ""),
        ("使用時機", ""),
    ], widths=(5.0, 11.5))

    add_heading(doc, "相對強弱描述", 2)
    add_body(doc, "如果還無法決定數字，可以先寫相對關係。這些描述之後能轉成第一版參數。")
    add_bullets(doc, [
        "這張卡比基礎資源卡昂貴，但可以重複使用。",
        "這張戰鬥卡前期難以製作，但能處理大量屍群。",
        "這張食物卡效率較低，優點是使用次數多。",
        "這張卡平時普通，但在低士氣階段特別有效。",
    ])
    add_prompt(doc, "這張卡最重要的取捨", 3)
    add_prompt(doc, "玩家會在什麼情況下選擇它", 3)

    add_page_break(doc)
    add_heading(doc, "數值與體驗目標", 1)
    add_body(doc, "離線時適合先決定數值要造成的體驗，再建立測試值。不要把暫定數字當成最終平衡。")
    add_two_col_table(doc, [
        ("起始資源", "玩家一開始大約能完成幾件事"),
        ("普通殭屍", "通常需要幾張戰鬥卡或多少戰鬥值"),
        ("忽視威脅", "放任一輪的代價應該多嚴重"),
        ("食物短缺", "連續幾輪後應進入危險"),
        ("士氣下降", "從 100 降到 70 約容許幾次失誤"),
        ("建築施工", "一棟基礎建築預計幾輪完成"),
        ("卡牌製作", "立即、單輪或多輪完成"),
        ("每輪產能", "玩家通常可推進幾項建設或製作"),
    ], header=("項目", "需要決定的體驗"))

    add_heading(doc, "體驗敘述範本", 2)
    add_checklist(doc, [
        "一次食物短缺會造成壓力，但不會立即結束遊戲。",
        "連續三輪沒有處理食物，應足以跨越一個士氣階段。",
        "普通殭屍可以臨時補救，屍群必須提前準備。",
        "第一棟建築應在第一輪或第二輪完成。",
    ])
    add_prompt(doc, "我希望玩家在前五輪感受到的壓力", 4)
    add_prompt(doc, "我最不希望出現的無聊或必勝策略", 3)

    add_page_break(doc)
    add_heading(doc, "事件內容", 1)
    add_body(doc, "事件的重點是讓玩家在有限資源下做出有代價的選擇。故事文字可以之後補寫，先確認每個選項改變了什麼。")
    add_two_col_table(doc, [
        ("事件名稱", ""),
        ("事件情境", ""),
        ("玩家為何要決定", ""),
        ("觸發條件", "日期、士氣、建築、殭屍或其他條件"),
    ], widths=(5.0, 11.5))

    add_heading(doc, "選項設計", 2)
    add_two_col_table(doc, [
        ("選項一", "付出、立即獲得、長期後果"),
        ("選項二", "付出、風險、長期後果"),
        ("選項三", "拒絕或保守選擇，以及其代價"),
        ("不可選原因", "資源不足、條件不足或其他限制"),
    ], widths=(4.5, 12.0))
    add_prompt(doc, "這個事件真正考驗玩家的能力", 3)
    add_prompt(doc, "是否存在明顯優於其他選項的答案", 3)

    add_page_break(doc)
    add_heading(doc, "士氣與失敗節奏", 1)
    add_body(doc, "士氣系統需要定義各階段的狀態、跨越門檻的代價與恢復方式。每個效果都應讓玩家改變決策，而不是只改一個數字。")
    add_two_col_table(doc, [
        ("Stage 0", "效果名稱、正面效果、失去時的感受"),
        ("Stage 1", "正面效果、負面效果、補救方式"),
        ("Stage 2", "正面效果、負面效果、最後補救窗口"),
        ("士氣歸零", "立即失敗、最後事件或其他結局"),
        ("永久上限損傷", "來源、頻率、是否存在有限恢復方式"),
    ], widths=(4.5, 12.0))
    add_checklist(doc, [
        "食物短缺是主要士氣壓力之一。",
        "殭屍存活到結算才造成損失。",
        "恢復士氣不會逆轉已失去的 Stage。",
        "低士氣效果應改變玩法，而非只有數值懲罰。",
    ])
    add_prompt(doc, "我希望玩家跨越門檻時感受到什麼", 3)
    add_prompt(doc, "低士氣玩法可以提供哪些特殊機會", 3)

    add_page_break(doc)
    add_heading(doc, "情報中心頁面", 1)
    add_body(doc, "情報中心可以收納需要查閱但不必常駐 HUD 的資料。每一頁都應定義何時公開，以及玩家能否操作。")
    add_two_col_table(doc, [
        ("基地目前狀態", "一開始可見"),
        ("食物計畫", "情報中心建成或指定等級後解鎖"),
        ("士氣金字塔紀錄", "依局前 Profile 或遊戲進度公開"),
        ("地區情報", "探索或情報來源解鎖"),
        ("殭屍類型", "遭遇後解鎖"),
        ("劇情線索與收藏", "事件發現後解鎖"),
        ("上一輪收支報告", "未來功能，尚未實作"),
    ], header=("候選頁面", "公開時機"))

    add_heading(doc, "新資訊頁資料卡", 2)
    add_two_col_table(doc, [
        ("頁面名稱", ""),
        ("公開條件", "一開始、建築等級、事件或永遠隱藏"),
        ("玩家要解決的問題", ""),
        ("顯示資料", ""),
        ("可進行的操作", "唯讀或可操作"),
    ], widths=(5.2, 11.3))

    add_page_break(doc)
    add_heading(doc, "下一步優先順序", 1)
    add_body(doc, "下一個開發項目應先驗證一種主要樂趣。請選擇目前最想確認的方向，並寫下選擇原因。")
    add_checklist(doc, [
        "建築經營與製作",
        "士氣下降與失敗閉環",
        "殭屍壓力與戰鬥準備",
        "探索與事件內容量",
        "卡牌構築與強化",
        "美術回饋與沉浸感",
    ])
    add_prompt(doc, "我選擇的下一個方向", 2)
    add_prompt(doc, "完成後要證明什麼", 3)
    add_prompt(doc, "哪些內容暫時不做", 3)

    add_heading(doc, "需要回到電腦前再決定的項目", 1)
    add_bullets(doc, [
        "UI 的精確尺寸、座標、字體大小與文字擁擠程度。",
        "動畫速度、停頓時間、拖曳與滑鼠命中手感。",
        "大量物件、粒子與動態光影的效能。",
        "最終資源平衡與玩家是否真正理解提示。",
    ])

    add_heading(doc, "本次離線決策摘要", 1)
    add_two_col_table(doc, [
        ("今天完成的決定", ""),
        ("仍有兩個候選方案", ""),
        ("需要進遊戲測試", ""),
        ("下一次要討論", ""),
    ], widths=(5.2, 11.3))

    # Normalize paragraph fonts, including list styles and table content.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                style_run(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name is None:
                            style_run(run)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "殭屍生存卡牌離線設計決策工作簿"
    doc.core_properties.subject = "離線整理建築 卡牌 數值 事件 士氣與開發優先順序"
    doc.core_properties.author = "Zombie Card Survive Project"
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    build_document()
